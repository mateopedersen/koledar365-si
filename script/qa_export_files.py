#!/usr/bin/env python3
"""Parse the actual downloads/print PDFs saved by qa_exports.mjs."""
import calendar,datetime,json,os
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
folder=Path(os.environ.get('QA_OUT','evidence/qa/exports'));report=json.loads((folder/'exports.json').read_text());rows=[]
for source in report['rows']:
 row={**source,'file_checks':[]}
 try:
  assert source['passed'],source.get('error','Browser export failed')
  file=folder/source['file']
  if source['kind']=='xlsx':
   wb=load_workbook(file);values=load_workbook(file,data_only=True)
   y,m=source['year'],source['month'];count=calendar.monthrange(y,m)[1] if m else 365+calendar.isleap(y)
   assert len(wb.worksheets)==(3 if m else 14)
   dates=wb['Datumi'];cached=values['Datumi'];first=datetime.date(y,m or 1,1)
   working=0
   for i in range(count):
    r=i+5;date=first+datetime.timedelta(days=i)
    assert dates[f'A{r}'].value.date()==date
    assert dates[f'F{r}'].data_type=='f' and 'WEEKDAY' in dates[f'F{r}'].value
    expected=int(date.weekday()<5 and dates[f'D{r}'].value=='Ne')
    assert cached[f'F{r}'].value==expected;working+=expected
   assert dates[f'F{count+5}'].value==f'=SUM(F5:F{count+4})'
   assert cached[f'F{count+5}'].value==working
   assert all(s.page_setup.orientation=='landscape' for s in wb.worksheets)
   row['file_checks']=['OOXML parsed','Slovene sheets','date cells','editable notes','workday formulas and cached results','A4 landscape setup']
  else:
   assert file.read_bytes().startswith(b'%PDF-')
   pdf=PdfReader(file,strict=True);assert len(pdf.pages)>0
   if source['kind']=='print' and source.get('name') in ['annual','monthly','school']:
    assert len(pdf.pages)==1,'Calendar should fit one sheet without a detached heading/footer'
   if source['kind']=='print' and source.get('name')=='lunar-year':
    first=pdf.pages[0].extract_text()
    assert all(month in first for month in ['Januar','Februar','Marec','April','Maj','Junij','Julij','Avgust','September','Oktober','November','December']),'All twelve lunar months must fit the first print sheet'
    assert all('2028' in page.extract_text() for page in pdf.pages),'A lunar print page contains only source notes'
   for page in pdf.pages:
    assert len(page.extract_text().strip())>30,'Blank or footer-only print page'
   expected=(297,420) if source['paper']=='A3' else (210,297)
   if source['orientation']=='landscape':expected=expected[::-1]
   for page in pdf.pages:
    got=(float(page.mediabox.width)*25.4/72,float(page.mediabox.height)*25.4/72)
    assert all(abs(a-b)<1.2 for a,b in zip(got,expected)),f'Expected {expected}, got {got}'
   content=' '.join(page.extract_text() for page in pdf.pages)
   if source['kind']=='print' and source.get('name')=='student':
    assert '7. 2. 2028' in content and '11. 2. 2028' in content,'Selected school week missing'
    assert 'Matematika' in content,'Saved sample lesson missing'
   if source['kind']=='personal-pdf':
    if source['path']=='/dnevni-planer/':
     assert all(f'{hour:02}:00' in content for hour in range(7,22)),'Missing daily hour after pagination'
    if source['path']=='/solski-urnik/':
     assert all(f'{hour}. ura' in content for hour in range(1,10)),'Missing lesson after pagination'
    if source['path']=='/rojstnodnevni-koledar/':
     assert all(month in content for month in ['Januar','Februar','Marec','April','Maj','Junij','Julij','Avgust','September','Oktober','November','December']),'Missing birthday month after pagination'
   row['pages']=len(pdf.pages);row['file_checks']=['PDF parsed','non-empty pages','selected paper size','selected orientation']
  row['file_passed']=True
 except Exception as e:row['file_passed']=False;row['file_error']=str(e)
 rows.append(row)
result={'tested':len(rows),'failed':sum(not r['file_passed'] for r in rows),'rows':rows}
(folder/'file-verification.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps({'tested':result['tested'],'failed':result['failed'],'failures':[r for r in rows if not r['file_passed']]}));raise SystemExit(bool(result['failed']))
